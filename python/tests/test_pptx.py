"""Checks of the python-pptx layer. Runs on CPython and on rpy:

    PYTHONPATH=python python3 python/tests/test_pptx.py
    rpy --mount python=/lib --path /lib python/tests/test_pptx.py
"""

import io
import json
import os
import sys

from pptx import Presentation
from pptx.chart.data import BubbleChartData, CategoryChartData, ChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Emu, Inches, Pt

import sliqtly
from pptx._markdown import d3_format


def fences(md, lang="vega-lite"):
    out = []
    parts = md.split("```" + lang + "\n")
    for p in parts[1:]:
        out.append(json.loads(p.split("\n```", 1)[0]))
    return out


def test_units():
    assert Inches(1) == 914400 and Pt(1) == 12700 and Cm(1) == 360000
    assert Inches(2).inches == 2.0 and Pt(18).pt == 18.0
    assert isinstance(Inches(1) + Inches(1), int)
    assert Emu(914400).inches == 1.0


def test_enums_and_colors():
    assert PP_ALIGN.CENTER == 2 and PP_ALIGN.CENTER.name == "CENTER"
    assert str(RGBColor(0x12, 0x34, 0xAB)) == "1234AB"
    assert RGBColor.from_string("FF0000") == (255, 0, 0)
    assert MSO_SHAPE.ROUNDED_RECTANGLE == 5


def test_object_model():
    prs = Presentation()
    assert prs.slide_width == Inches(10) and len(prs.slide_layouts) == 11
    assert prs.slide_layouts.get_by_name("Title Only") is prs.slide_layouts[5]
    s = prs.slides.add_slide(prs.slide_layouts[1])
    assert s.slide_id == 256 and prs.slides.index(s) == 0
    assert s.shapes.title.placeholder_format.type == PP_PLACEHOLDER.TITLE
    assert [p.placeholder_format.idx for p in s.placeholders] == [0, 1]
    tb = s.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
    assert tb.shape_type == MSO_SHAPE_TYPE.TEXT_BOX and tb.name == "TextBox 3"
    tb.text = "a\nb"
    assert [p.text for p in tb.text_frame.paragraphs] == ["a", "b"]
    p = tb.text_frame.paragraphs[0]
    p.text = "x\ny"
    assert p.text == "x\vy"
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, 0, 0, Inches(1), Inches(1))
    shp.fill.solid()
    shp.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
    assert shp.auto_shape_type == MSO_SHAPE.OVAL and len(s.shapes) == 4
    t = s.shapes.add_table(2, 2, 0, 0, Inches(4), Inches(1)).table
    t.columns[0].width = Inches(3)
    t.cell(0, 0).merge(t.cell(0, 1))
    assert t.cell(0, 0).is_merge_origin and t.cell(0, 1).is_spanned and t.cell(0, 0).span_width == 2
    # the python-pptx way of deleting a slide
    prs.slides.add_slide(prs.slide_layouts[6])
    ids = prs.slides._sldIdLst
    ids.remove(list(ids)[0])
    assert len(prs.slides) == 1 and prs.slides[0].slide_layout.name == "Blank"


def test_markdown_text():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Deck {draft}"
    s.placeholders[1].text = "Subtitle"
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "List"
    tf = s.placeholders[1].text_frame
    tf.text = "One"
    p = tf.add_paragraph()
    p.text = "Two"
    p.level = 1
    r = tf.add_paragraph().add_run()
    r.text = "Link"
    r.hyperlink.address = "https://example.com"
    s.notes_slide.notes_text_frame.text = "Say this."
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    tb.text = "Big title"
    tb.text_frame.paragraphs[0].font.size = Pt(36)
    body = s.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(1))
    body.text = "Costs $5 *each* <b> 1. item"
    body.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(0, Inches(3), Inches(2), Inches(1)).text = "plain"
    md = prs.to_markdown()
    assert "# Deck \\{draft} \n" not in md and "# Deck \\{draft}\n" in md
    assert "\nSubtitle\n" in md
    assert "- One\n  - Two\n- [Link](https://example.com)" in md
    assert "::: notes\nSay this.\n:::" in md
    assert "## Big title\n" in md
    assert "Costs \\$5 \\*each\\* &lt;b> 1. item\n{.center}" in md
    assert "## Slide 4 {heading=hidden}" in md


def test_markdown_layout():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Side by side"
    s.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(6), Inches(4)).text = "Left"
    s.shapes.add_textbox(Inches(7), Inches(1.5), Inches(2.5), Inches(4)).text = "Right"
    s.shapes.add_textbox(Inches(0.5), Inches(6), Inches(9), Inches(1)).text = "Below"
    s = prs.slides.add_slide(prs.slide_layouts[3])
    s.shapes.title.text = "Two content"
    s.placeholders[1].text = "A"
    s.placeholders[2].text = "B"
    md = prs.to_markdown()
    assert ":::: columns\n::: col\nLeft\n:::\n::: col\nRight\n:::\n::::\n{widths=\"71 29\"}\n\nBelow" in md, md
    assert ":::: columns\n::: col\n- A\n:::\n::: col\n- B\n:::\n::::" in md


def test_markdown_shapes_and_tables():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    card = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(1))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0x11, 0x22, 0x33)
    card.text = "42%\nMargin"
    sliqtly.animate(card).zoom(0.4)
    s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5), Inches(1), Inches(1), Inches(1))
    t = s.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    t.cell(0, 0).text = "A|B"
    t.cell(0, 1).text = "C"
    t.cell(1, 0).text = "1"
    t.cell(1, 1).text = "2"
    t2 = s.shapes.add_table(2, 2, Inches(1), Inches(5), Inches(4), Inches(1)).table
    t2.cell(0, 0).merge(t2.cell(0, 1))
    t2.cell(0, 0).text = "Both"
    t2.cell(1, 0).fill.solid()
    t2.cell(1, 0).fill.fore_color.rgb = RGBColor(255, 0, 0)
    md = prs.to_markdown()
    assert "42%<br>Margin\n{container=box background=#112233 anim=zoom seconds=0.4}" in md, md
    assert "| A\\|B | C |\n| --- | --- |\n| 1 | 2 |" in md
    assert '<th colspan="2">Both</th>' in md and '<td style="background:#FF0000"></td>' in md
    assert any("Oval" in w for w in prs.warnings)


def test_charts():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    cd = CategoryChartData(number_format="0%")
    cd.categories = ["a", "b"]
    cd.add_series("S1", (0.1, 0.2))
    cd.add_series("S2", (0.3, None))
    ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, 0, 0, Inches(4), Inches(3), cd).chart
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.plots[0].has_data_labels = True
    ch.series[0].format.fill.solid()
    ch.series[0].format.fill.fore_color.rgb = RGBColor(1, 2, 3)
    ch.series[1].format.fill.solid()
    ch.series[1].format.fill.fore_color.rgb = RGBColor(4, 5, 6)
    one = ChartData()
    one.categories = ["x", "y"]
    one.add_series("Only", (1, 2))
    s.shapes.add_chart(XL_CHART_TYPE.PIE, 0, Inches(4), Inches(4), Inches(3), one)
    xy = XyChartData()
    xs = xy.add_series("pts")
    xs.add_data_point(1, 2)
    s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(5), 0, Inches(4), Inches(3), xy)
    bub = BubbleChartData()
    bub.add_series("b").add_data_point(1, 2, 3)
    s.shapes.add_chart(XL_CHART_TYPE.BUBBLE, Inches(5), Inches(4), Inches(4), Inches(3), bub)
    specs = fences(prs.to_markdown())
    bar, pie, scatter, bubble = specs
    assert bar["encoding"]["y"]["type"] == "nominal" and bar["encoding"]["x"]["stack"] == "zero"
    assert bar["encoding"]["color"]["scale"]["range"] == ["#010203", "#040506"]
    assert bar["encoding"]["color"]["legend"] == {"orient": "top"}
    assert len(bar["data"]["values"]) == 3 and bar["layer"][1]["encoding"]["text"]["format"] == ".0%"
    assert pie["title"] == "Only" and pie["mark"]["type"] == "arc"
    assert scatter["data"]["values"] == [{"x": 1, "y": 2}]
    assert bubble["encoding"]["size"]["field"] == "size"


def test_formats():
    assert d3_format("General") is None
    assert d3_format("0%") == ".0%" and d3_format("0.0%") == ".1%"
    assert d3_format("#,##0") == ",.0f" and d3_format("#,##0.00") == ",.2f"
    assert d3_format('"$"#,##0') == ",.0f" and d3_format("$#,##0.0") == "$,.1f"


def test_sliqtly_blocks_and_save():
    prs = Presentation()
    sliqtly.deck(prs, lang="fi", footer_right="{page}")
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Blocks"
    sliqtly.transition(s, "zoom", 0.5)
    s.shapes.add_vega({"mark": "bar", "data": {"values": [{"a": 1}]}})
    s.shapes.add_mermaid("flowchart LR\n  A --> B")
    s.shapes.add_markdown("> quote")
    png = b"\x89PNG\r\n\x1a\n" + b"\0\0\0\rIHDR" + (64).to_bytes(4, "big") + (32).to_bytes(4, "big") + b"\0" * 20
    folder = "/tmp/pptx-test"
    os.makedirs(folder, exist_ok=True)
    with open(folder + "/logo.png", "wb") as f:
        f.write(png)
    pic = s.shapes.add_picture(folder + "/logo.png", 0, Inches(6), width=Inches(2))
    assert pic.height == Inches(1)
    path = prs.save(folder + "/deck.pptx")
    assert path == folder + "/deck.md"
    md = open(path).read()
    assert md.startswith('---\nlang: fi\nfooter-right: "{page}"\n---\n')
    assert "## Blocks {transition=zoom seconds=0.5}" in md
    assert fences(md)[0]["background"] == "rgba(0,0,0,0)"
    assert "```mermaid\nflowchart LR\n  A --> B\n```" in md and "> quote" in md
    assert "![logo](media/logo.png)" in md
    assert open(folder + "/media/logo.png", "rb").read() == png
    buf = io.BytesIO()
    prs.save(buf)
    assert buf.getvalue().decode("utf-8") == md


if __name__ == "__main__":
    failed = 0
    tests = [(k, v) for k, v in list(globals().items()) if k.startswith("test_")]
    for name, fn in tests:
        try:
            fn()
            print("ok  ", name)
        except Exception as e:
            failed += 1
            print("FAIL", name, type(e).__name__, e)
    print("%d/%d passed" % (len(tests) - failed, len(tests)))
    sys.exit(1 if failed else 0)
